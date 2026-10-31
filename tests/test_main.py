"""Regression tests for the repository-root Worklog package."""
import os
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PARENT = ROOT.parent
if str(PARENT) not in sys.path:
    sys.path.insert(0, str(PARENT))

from Worklog.__main__ import (  # noqa: E402
    convert_duration,
    find_first_empty_row,
    get_excel_path,
    is_duplicate,
    parse_line,
    process_data,
    split_datetime,
)


def test_default_excel_file():
    path, name = get_excel_path()
    assert Path(path).resolve() == ROOT / "Worklog.xlsx"
    assert name == "Worklog"


def test_parse_line():
    result = parse_line('{"开发":"2026年7月29日 17:45,2:30:00"}')
    assert result == ("开发", "2026-07-29", "17:45", 2.5)


def test_bad_input():
    assert parse_line("not json") is None
    assert parse_line('{"": "2026年7月29日 17:45,1:00:00"}') is None
    assert parse_line("") is None


def test_datetime_and_duration():
    assert split_datetime("2026年7月9日 7:05") == ("2026-07-09", "07:05")
    assert convert_duration("2:00:00") == 2
    assert convert_duration("1:30:00") == 1.5
    assert convert_duration("invalid") == 0


def test_process_data_and_duplicate_prevention(tmp_path):
    target = tmp_path / "Worklog.xlsx"
    line = '{"开发":"2026年7月29日 17:45,2:30:00"}'
    process_data([line], str(target), "Worklog")
    process_data([line], str(target), "Worklog")
    wb = load_workbook(target)
    ws = wb["Worklog"]
    assert ws.max_row == 2
    assert ws["A2"].value == "开发"
    assert ws["D2"].value == 2.5
    assert is_duplicate(ws, ("开发", "2026-07-29", "17:45", 2.5))
    assert find_first_empty_row(ws) == 3
    wb.close()


def test_module_cli_exits_from_stdin():
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONPATH"] = str(PARENT)
    result = subprocess.run(
        [sys.executable, "-m", "Worklog"],
        input="exit\n",
        cwd=PARENT,
        env=env,
        text=True,
        capture_output=True,
        encoding="utf-8",
        timeout=15,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert "Worklog.xlsx" in result.stdout
